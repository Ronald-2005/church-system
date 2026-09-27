from django.shortcuts import render
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse
from openpyxl import Workbook
import openpyxl
from members.models import Member  # assuming your members app model is Member
from contributions.models import Contribution  # assuming your contributions app model is Contribution
from decimal import Decimal

def is_admin(user):
    return user.is_superuser or user.is_staff


@login_required
def membership_report(request):
    members = Member.objects.all()

    total_members = members.count()
    male_members = members.filter(gender="Male").count()
    female_members = members.filter(gender="Female").count()

    context = {
        "members": members,
        "total_members": total_members,
        "male_members": male_members,
        "female_members": female_members,
    }

    return render(request, "reports/membership_report.html", context)


@login_required
@user_passes_test(is_admin)
def download_membership_report(request):
    members = Member.objects.all()

    # Create Excel workbook
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Membership Report"

    # Headers
    sheet.append(["ID", "Full Name", "Gender", "Phone", "Address", "Date of Birth"])

    # Data rows
    for member in members:
        sheet.append([
            member.id,
            member.full_name,
            member.gender,
            member.phone,
            member.address,
            member.date_of_birth.strftime("%Y-%m-%d") if member.date_of_birth else "",
        ])

    # Prepare response
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="membership_report.xlsx"'

    wb.save(response)
    return response


# Create your views here.

@login_required
def contribution_report(request):
    contributions = Contribution.objects.all()
    return render(request, 'reports/contribution_report.html', {'contributions': contributions})


@login_required

def download_contribution_report(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contribution Report"

    # Header row
    sheet.append(["ID", "Giver", "Amount", "Payment Method", "Date", "M-Pesa Code", "Phone", "Status", "Notes"])

    total_amount = Decimal('0.00')

    # Data rows
    for c in Contribution.objects.all():
        sheet.append([
            c.id,
            getattr(c.giver, 'full_name', ''),
            c.amount,
            c.payment_method,
            c.date.strftime("%Y-%m-%d %H:%M"),
            c.mpesa_transaction_code or '',
            c.mpesa_phone_number or '',
            c.mpesa_status or '',
            c.notes or ''
        ])
        total_amount += c.amount

    # Add a blank row before total
    sheet.append([])
    # Add total row
    sheet.append(["", "TOTAL CONTRIBUTIONS", total_amount, "", "", "", "", "", ""])

    # Style: bold header and total (optional)
    from openpyxl.styles import Font
    bold_font = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold_font
    for cell in sheet[sheet.max_row]:
        cell.font = bold_font

    # Response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="contribution_report.xlsx"'
    workbook.save(response)
    return response

from django.shortcuts import render
from django.contrib.auth.decorators import login_required

@login_required
def report_home(request):
    return render(request, 'reports/report_home.html')
